# 验证 PDF、DOCX、TXT 简历解析和文件安全限制。

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.core.exceptions import ResumeFileError
from app.services import resume_file_parser
from app.services.resume_file_parser import (
    parse_resume_file,
)


def make_txt_bytes(text: str) -> bytes:
    return text.encode("utf-8")


def make_pdf_bytes(text: str, *, pages: int = 1) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for page_index in range(pages):
        if page_index == 0:
            pdf.drawString(72, 720, text)
        else:
            pdf.drawString(72, 720, f"page {page_index + 1}")
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("技能", level=1)
    document.add_paragraph("Python、FastAPI、PostgreSQL")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "实现 RAG 混合检索"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_txt_resume() -> None:
    result = parse_resume_file(
        filename="resume.txt",
        content_type="text/plain",
        data=make_txt_bytes("Python 后端开发经验"),
    )

    assert result.source_type == "txt"
    assert result.char_count > 0
    assert "Python" in result.text


def test_parse_docx_resume_with_table() -> None:
    result = parse_resume_file(
        filename="resume.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        data=make_docx_bytes(),
    )

    assert result.source_type == "docx"
    assert "Python" in result.text
    assert "实现 RAG 混合检索" in result.text


def test_parse_text_pdf_resume() -> None:
    result = parse_resume_file(
        filename="resume.pdf",
        content_type="application/pdf",
        data=make_pdf_bytes("FastAPI backend resume"),
    )

    assert result.source_type == "pdf"
    assert result.page_count == 1
    assert "FastAPI" in result.text


def test_legacy_doc_is_rejected() -> None:
    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.doc",
            content_type="application/msword",
            data=b"not a docx",
        )

    assert exc_info.value.error_code == "unsupported_file_type"


def test_pdf_without_text_returns_ocr_error() -> None:
    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="scan.pdf",
            content_type="application/pdf",
            data=make_pdf_bytes(""),
        )

    assert exc_info.value.error_code == "scanned_pdf_requires_ocr"


def test_encrypted_pdf_is_rejected() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("secret")
    writer.write(buffer)

    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="encrypted.pdf",
            content_type="application/pdf",
            data=buffer.getvalue(),
        )

    assert exc_info.value.error_code == "encrypted_pdf"


def test_file_size_limit(monkeypatch) -> None:
    monkeypatch.setattr(resume_file_parser, "MAX_RESUME_FILE_BYTES", 5)

    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.txt",
            content_type="text/plain",
            data=b"123456",
        )

    assert exc_info.value.error_code == "file_too_large"


def test_text_length_limit(monkeypatch) -> None:
    monkeypatch.setattr(resume_file_parser, "MAX_RESUME_TEXT_LENGTH", 5)

    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.txt",
            content_type="text/plain",
            data=b"123456",
        )

    assert exc_info.value.error_code == "resume_text_too_long"


def test_docx_zip_bomb_limit(monkeypatch) -> None:
    monkeypatch.setattr(
        resume_file_parser,
        "MAX_DOCX_UNCOMPRESSED_BYTES",
        1,
    )

    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            data=make_docx_bytes(),
        )

    assert exc_info.value.error_code == "docx_zip_bomb"


def test_pdf_page_limit(monkeypatch) -> None:
    monkeypatch.setattr(resume_file_parser, "MAX_PDF_PAGES", 1)

    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.pdf",
            content_type="application/pdf",
            data=make_pdf_bytes("resume", pages=2),
        )

    assert exc_info.value.error_code == "pdf_too_many_pages"


def test_empty_file_is_rejected() -> None:
    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.txt",
            content_type="text/plain",
            data=b"",
        )

    assert exc_info.value.error_code == "empty_file"


def test_fake_docx_is_rejected() -> None:
    with pytest.raises(ResumeFileError) as exc_info:
        parse_resume_file(
            filename="resume.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            data=b"PK\x03\x04not-a-real-docx",
        )

    assert exc_info.value.error_code == "invalid_docx"
