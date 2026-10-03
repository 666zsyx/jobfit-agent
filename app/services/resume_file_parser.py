# 解析 PDF / DOCX / TXT 简历文件，并执行安全限制。

from __future__ import annotations

import asyncio
import io
import logging
import zipfile
from pathlib import PurePath

import pdfplumber
from docx import Document
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.exceptions import ResumeFileError


logger = logging.getLogger(__name__)

MAX_RESUME_FILE_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 30
MAX_RESUME_TEXT_LENGTH = 50_000
MAX_DOCX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MIN_PDF_TEXT_LENGTH = 20

SOURCE_TYPE_BY_SUFFIX = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
}
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "application/octet-stream",
}


class ResumeParseResult(BaseModel):
    text: str = Field(min_length=1, max_length=MAX_RESUME_TEXT_LENGTH)
    source_type: str
    original_filename: str
    page_count: int | None = None
    char_count: int = Field(ge=1, le=MAX_RESUME_TEXT_LENGTH)


def _file_error(
    message: str,
    error_code: str,
    *,
    status_code: int = 422,
) -> ResumeFileError:
    return ResumeFileError(
        message=message,
        error_code=error_code,
        status_code=status_code,
    )


def _safe_filename(filename: str | None) -> str:
    value = (filename or "resume").strip()
    return PurePath(value).name[:255] or "resume"


def _detect_source_type(
    *,
    filename: str,
    content_type: str | None,
    data: bytes,
) -> str:
    suffix = PurePath(filename).suffix.lower()
    normalized_content_type = (content_type or "").split(";")[0].strip().lower()

    if suffix == ".doc":
        raise _file_error(
            "暂不支持旧版 .doc，请转换为 DOCX",
            "unsupported_file_type",
            status_code=415,
        )

    if suffix not in SOURCE_TYPE_BY_SUFFIX:
        raise _file_error(
            "只支持 PDF、DOCX 和 TXT 文件",
            "unsupported_file_type",
            status_code=415,
        )

    if normalized_content_type and normalized_content_type not in ALLOWED_CONTENT_TYPES:
        raise _file_error(
            "文件 Content-Type 不受支持",
            "unsupported_file_type",
            status_code=415,
        )

    source_type = SOURCE_TYPE_BY_SUFFIX[suffix]
    if source_type == "pdf" and not data.startswith(b"%PDF-"):
        raise _file_error("文件不是有效 PDF", "corrupted_pdf")
    if source_type == "docx" and not data.startswith(b"PK\x03\x04"):
        raise _file_error("文件不是有效 DOCX", "invalid_docx")
    if source_type == "txt" and b"\x00" in data[:1024]:
        raise _file_error("TXT 文件包含二进制内容", "unsupported_file_type")

    return source_type


def _normalize_text(text: str) -> str:
    normalized = text.replace("\x00", "")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in normalized.split("\n")]
    compact: list[str] = []
    blank_count = 0
    for line in lines:
        if line.strip():
            blank_count = 0
            compact.append(line)
        else:
            blank_count += 1
            if blank_count <= 1:
                compact.append("")
    return "\n".join(compact).strip()


def _validate_text(text: str, *, error_code: str = "empty_resume_text") -> str:
    normalized = _normalize_text(text)
    if not normalized:
        raise _file_error("未能从文件中提取到简历文本", error_code)
    if len(normalized) > MAX_RESUME_TEXT_LENGTH:
        raise _file_error(
            "简历文本超过 50,000 字符限制",
            "resume_text_too_long",
        )
    return normalized


def _extract_pdf_with_pdfplumber(data: bytes) -> str:
    texts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text:
                texts.append(page_text)

            for table in page.extract_tables() or []:
                rows = [
                    " | ".join(
                        (cell or "").strip()
                        for cell in row
                    )
                    for row in table
                ]
                table_text = "\n".join(
                    row for row in rows if row.strip()
                )
                if table_text and table_text not in page_text:
                    texts.append(table_text)
    return "\n\n".join(texts)


def _extract_pdf_with_pypdf(reader: PdfReader) -> str:
    texts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text:
            texts.append(page_text)
    return "\n\n".join(texts)


def _parse_pdf(data: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise _file_error("PDF 文件损坏或无法读取", "corrupted_pdf") from exc

    if reader.is_encrypted:
        try:
            decrypt_result = reader.decrypt("")
        except Exception as exc:
            raise _file_error("PDF 已加密，无法解析", "encrypted_pdf") from exc
        if decrypt_result == 0:
            raise _file_error("PDF 已加密，无法解析", "encrypted_pdf")

    page_count = len(reader.pages)
    if page_count > MAX_PDF_PAGES:
        raise _file_error(
            "PDF 页数超过 30 页限制",
            "pdf_too_many_pages",
        )

    text = ""
    try:
        text = _extract_pdf_with_pdfplumber(data)
    except Exception:
        logger.warning("pdfplumber 解析失败，尝试 pypdf 回退", exc_info=True)

    if len(_normalize_text(text)) < MIN_PDF_TEXT_LENGTH:
        fallback_text = _extract_pdf_with_pypdf(reader)
        if len(_normalize_text(fallback_text)) > len(_normalize_text(text)):
            text = fallback_text

    normalized = _normalize_text(text)
    if len(normalized) < MIN_PDF_TEXT_LENGTH:
        raise _file_error(
            "PDF 没有可提取文本，可能是扫描版，需要 OCR",
            "scanned_pdf_requires_ocr",
        )
    return _validate_text(normalized), page_count


def _validate_docx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            total_size = sum(info.file_size for info in infos)
            if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise _file_error(
                    "DOCX 解压后过大",
                    "docx_zip_bomb",
                )
            for info in infos:
                normalized_name = info.filename.replace("\\", "/")
                if normalized_name.startswith("/") or "../" in normalized_name:
                    raise _file_error(
                        "DOCX 包含非法路径",
                        "invalid_docx",
                    )
    except ResumeFileError:
        raise
    except zipfile.BadZipFile as exc:
        raise _file_error("DOCX 文件损坏", "invalid_docx") from exc


def _parse_docx(data: bytes) -> str:
    _validate_docx_archive(data)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise _file_error("DOCX 文件损坏或无法读取", "invalid_docx") from exc

    texts: list[str] = []
    texts.extend(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.text.strip()
    )
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text and cell.text.strip()
            ]
            if cells:
                texts.append(" | ".join(cells))
    return _validate_text("\n".join(texts))


def _decode_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise _file_error("TXT 文件编码无法识别", "invalid_text_encoding")


def parse_resume_file(
    *,
    filename: str | None,
    content_type: str | None,
    data: bytes,
) -> ResumeParseResult:
    """同步解析简历文件；调用方应在线程中执行。"""

    if not data:
        raise _file_error("文件为空", "empty_file", status_code=400)
    if len(data) > MAX_RESUME_FILE_BYTES:
        raise _file_error("文件超过 10 MB 限制", "file_too_large", status_code=413)

    safe_filename = _safe_filename(filename)
    source_type = _detect_source_type(
        filename=safe_filename,
        content_type=content_type,
        data=data,
    )

    page_count: int | None = None
    if source_type == "pdf":
        text, page_count = _parse_pdf(data)
    elif source_type == "docx":
        text = _parse_docx(data)
    else:
        text = _validate_text(_decode_txt(data))

    return ResumeParseResult(
        text=text,
        source_type=source_type,
        original_filename=safe_filename,
        page_count=page_count,
        char_count=len(text),
    )


async def parse_resume_file_async(
    *,
    filename: str | None,
    content_type: str | None,
    data: bytes,
) -> ResumeParseResult:
    return await asyncio.to_thread(
        parse_resume_file,
        filename=filename,
        content_type=content_type,
        data=data,
    )


__all__ = [
    "MAX_DOCX_UNCOMPRESSED_BYTES",
    "MAX_PDF_PAGES",
    "MAX_RESUME_FILE_BYTES",
    "MAX_RESUME_TEXT_LENGTH",
    "ResumeParseResult",
    "parse_resume_file",
    "parse_resume_file_async",
]
